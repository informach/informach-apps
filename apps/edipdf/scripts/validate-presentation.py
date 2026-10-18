"""Independent check of the synthetic presentation outputs, not an Office visual certification."""
from pathlib import Path
import sys
ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / '.cache/edipdf-analysis/python-odf'))
from pptx import Presentation
from odf.opendocument import load
from odf import draw, text, teletype
directory = Path(sys.argv[1])
presentation = Presentation(directory / 'exported.pptx')
assert len(presentation.slides) == 2
assert presentation.slide_width == 720 * 12700
assert presentation.slide_height == 405 * 12700
first = presentation.slides[0]
values = ' '.join(shape.text for shape in first.shapes if shape.has_text_frame)
assert 'AÇÃO_LOCAL_αβγ' in values and 'café e informação' in values, values
title = next(shape for shape in first.shapes if shape.has_text_frame and 'AÇÃO_LOCAL' in shape.text)
assert title.text_frame.paragraphs[0].runs[0].font.bold is True
assert title.text_frame.paragraphs[0].runs[0].font.size.pt == 24
assert sum(1 for shape in first.shapes if shape.shape_type == 13) == 1
assert any(shape.has_text_frame and 'ULTIMO_SLIDE' in shape.text for shape in presentation.slides[1].shapes)
document = load(str(directory / 'exported.odp'))
pages = document.body.getElementsByType(draw.Page)
assert len(pages) == 2
values = ' '.join(teletype.extractText(paragraph) for paragraph in pages[0].getElementsByType(text.P))
assert 'AÇÃO_LOCAL_αβγ' in values and 'café e informação' in values, values
assert len(pages[0].getElementsByType(draw.Image)) == 1
assert len(pages[0].getElementsByType(draw.Rect)) == 1
assert len(pages[0].getElementsByType(draw.Ellipse)) == 1
assert len(pages[0].getElementsByType(draw.Line)) == 1
assert 'ULTIMO_SLIDE' in teletype.extractText(pages[1])
print('PPTX/ODP: independent slide sizes/order, editable Unicode text/style, image and native shapes passed')
