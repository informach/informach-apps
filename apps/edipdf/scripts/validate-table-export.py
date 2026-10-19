"""Independent output check; isolated openpyxl 3.1.5, odfpy 1.4.1, defusedxml 0.7.1.
Runs only on the synthetic outputs from verify-table-export.mts, never production data.
"""
from pathlib import Path
import csv
import hashlib
import json
import sys

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / '.cache/edipdf-analysis/python-odf'))
from openpyxl import load_workbook
from odf.opendocument import load
from odf import table, text, teletype

directory = Path(sys.argv[1])
expected = [['Produto', 'Valor'], ['Ação αβγ', '12,50'], ['=1+2', '']]
workbook = load_workbook(directory / 'table.xlsx', data_only=False, read_only=True)
assert workbook.sheetnames == ['Tabela', 'Origem das células', 'Revisão PDF']
values = [[cell if cell is not None else '' for cell in row] for row in workbook['Tabela'].iter_rows(values_only=True)]
assert values == expected, values
assert workbook['Tabela']['A3'].data_type == 's'
assert workbook['Origem das células']['C2'].value == 1
workbook.close()
document = load(str(directory / 'table.ods'))
sheets = document.spreadsheet.getElementsByType(table.Table)
assert len(sheets) == 3
values = [[teletype.extractText(cell) for cell in row.getElementsByType(table.TableCell)] for row in sheets[0].getElementsByType(table.TableRow)]
assert values == expected, values
with (directory / 'table.csv').open(encoding='utf-8-sig', newline='') as source:
    values = list(csv.reader(source, delimiter=';'))
assert values == [expected[0], expected[1], ["'=1+2", '']], values
evidence = json.loads((directory / 'table.json').read_text(encoding='utf-8'))
assert [[cell['value'] for cell in row] for row in evidence['rows']] == expected
assert evidence['source']['sha256'] == hashlib.sha256((directory / 'source.pdf').read_bytes()).hexdigest()
assert evidence['rows'][1][0]['address'] == 'R2C1'
assert evidence['rows'][1][0]['box'] is not None
assert evidence['rows'][2][1]['box'] is None
print('XLSX/ODS/CSV/JSON: independent Unicode, literal formula text, values and provenance passed')
