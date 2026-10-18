"""Developer-only fixtures authored independently with python-pptx 1.0.2 and Pillow 12.3.0."""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile
import hashlib
import json
import sys
from xml.etree import ElementTree
ROOT = Path(__file__).resolve().parents[4]
sys.path.insert(0, str(ROOT / '.cache/edipdf-analysis/python-odf'))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from PIL import Image
OUT = ROOT / 'apps/edipdf/tests/fixtures/presentation'
OUT.mkdir(parents=True, exist_ok=True)
presentation = Presentation()
presentation.slide_width = Inches(10)
presentation.slide_height = Inches(5.625)
slide = presentation.slides.add_slide(presentation.slide_layouts[6])
title = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(0.6))
run = title.text_frame.paragraphs[0].add_run()
run.text = 'PPTX_LOCAL_αβγ — Ação e informação'
run.font.size = Pt(24)
run.font.bold = True
run.font.color.rgb = RGBColor(38, 74, 115)
body = slide.shapes.add_textbox(Inches(0.5), Inches(1.2), Inches(4), Inches(1))
run = body.text_frame.paragraphs[0].add_run()
run.text = 'João e café; texto editável.'
run.font.size = Pt(16)
run.font.italic = True
run.hyperlink.address = 'https://example.invalid/not-fetched'
rectangle = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.5), Inches(2.5), Inches(2), Inches(1))
rectangle.name = 'THEME_RECTANGLE'
ellipse = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.75), Inches(2.5), Inches(1), Inches(1))
ellipse.fill.solid()
ellipse.fill.fore_color.rgb = RGBColor(255, 128, 0)
ellipse.line.color.rgb = RGBColor(18, 52, 86)
image = Image.new('RGB', (20, 10), (255, 0, 0))
for y in range(10):
    for x in range(10, 20):
        image.putpixel((x, y), (0, 255, 0))
data = BytesIO()
image.save(data, format='PNG')
picture = slide.shapes.add_picture(BytesIO(data.getvalue()), Inches(0.5), Inches(3), Inches(2), Inches(1))
picture.crop_right = 0.5
hidden = slide.shapes.add_textbox(Inches(4), Inches(4), Inches(3), Inches(1))
hidden.text = 'HIDDEN_SOURCE'
hidden._element.nvSpPr.cNvPr.set('hidden', '1')
last = presentation.slides.add_slide(presentation.slide_layouts[6])
last.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(1)).text = 'ULTIMO_SLIDE_INDEPENDENTE'
path = OUT / 'independent.pptx'
presentation.save(OUT / 'unsupported-gradient.pptx')
rectangle.fill.solid()
rectangle.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
rectangle.shadow.inherit = False
ellipse.shadow.inherit = False
presentation.save(path)
with ZipFile(path) as archive:
    theme = ElementTree.fromstring(archive.read('ppt/theme/theme1.xml'))
    color = theme.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}clrScheme/{http://schemas.openxmlformats.org/drawingml/2006/main}accent1/{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr').attrib['val']
(OUT / 'manifest.json').write_text(json.dumps({'generator': 'python-pptx 1.0.2 / Pillow 12.3.0', 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'themeRectangleColor': '#' + color, 'cropVisibleRgb': [255, 0, 0]}, indent=2) + '\n', encoding='utf-8')
print('Synthetic independent PPTX fixture generated')

# A second producer exercises ODP import independently from EdiPDF's writer.
from odf.opendocument import OpenDocumentPresentation
from odf import draw as odraw, text as otext, style as ostyle
odp = OpenDocumentPresentation()
layout = ostyle.PageLayout(name='LocalLayout')
layout.addElement(ostyle.PageLayoutProperties(pagewidth='720pt', pageheight='405pt', margin='0pt'))
odp.automaticstyles.addElement(layout)
master = ostyle.MasterPage(name='LocalMaster', pagelayoutname=layout)
odp.masterstyles.addElement(master)
title_style = ostyle.Style(name='TitleLocal', family='paragraph')
title_style.addElement(ostyle.TextProperties(fontsize='24pt', fontweight='bold', color='#264a73'))
odp.styles.addElement(title_style)
frame_style = ostyle.Style(name='TextFrame', family='graphic')
frame_style.addElement(ostyle.GraphicProperties(fill='none', stroke='none', padding='0pt'))
odp.styles.addElement(frame_style)
rectangle_style = ostyle.Style(name='RectangleLocal', family='graphic')
rectangle_style.addElement(ostyle.GraphicProperties(fill='solid', fillcolor='#4f81bd', stroke='solid', strokecolor='#123456', strokewidth='1pt'))
odp.styles.addElement(rectangle_style)
picture_style = ostyle.Style(name='PictureLocal', family='graphic')
picture_style.addElement(ostyle.GraphicProperties(imageopacity='100%', gamma='100%', luminance='0%', contrast='0%', colormode='standard', mirror='none'))
odp.styles.addElement(picture_style)
page = odraw.Page(name='Primeiro', masterpagename=master)
frame = odraw.Frame(stylename=frame_style, x='36pt', y='30pt', width='648pt', height='60pt')
box = odraw.TextBox()
box.addElement(otext.P(stylename=title_style, text='ODP_LOCAL_αβγ — Ação e informação'))
frame.addElement(box)
page.addElement(frame)
page.addElement(odraw.Rect(stylename=rectangle_style, x='400pt', y='180pt', width='200pt', height='80pt'))
visible = BytesIO()
Image.new('RGB', (10, 10), (255, 0, 0)).save(visible, format='PNG')
picture_path = odp.addPicture('Pictures/red.png', mediatype='image/png', content=visible.getvalue())
frame = odraw.Frame(stylename=picture_style, x='36pt', y='150pt', width='100pt', height='100pt')
frame.addElement(odraw.Image(href=picture_path, type='simple', show='embed', actuate='onLoad'))
page.addElement(frame)
odp.presentation.addElement(page)
page = odraw.Page(name='Segundo', masterpagename=master)
frame = odraw.Frame(stylename=frame_style, x='36pt', y='30pt', width='648pt', height='80pt')
box = odraw.TextBox()
box.addElement(otext.P(text='ULTIMO_SLIDE_ODP_INDEPENDENTE'))
frame.addElement(box)
page.addElement(frame)
odp.presentation.addElement(page)
odp.save(str(OUT / 'independent.odp'))
manifest = json.loads((OUT / 'manifest.json').read_text(encoding='utf-8'))
manifest['odpGenerator'] = 'odfpy 1.4.1 / Pillow 12.3.0'
manifest['odpSha256'] = hashlib.sha256((OUT / 'independent.odp').read_bytes()).hexdigest()
(OUT / 'manifest.json').write_text(json.dumps(manifest, indent=2) + '\n', encoding='utf-8')
print('Synthetic independent ODP fixture generated')
